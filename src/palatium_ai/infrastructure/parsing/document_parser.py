# src/palatium_ai/infrastructure/parsing/document_parser.py

"""DocumentParserPort adapters: plain text, PDF, DOCX, XLSX, PPTX, HTML.

Parsing is the first place hostile bytes are interpreted, so every adapter:

* runs its third-party library in a worker thread — the libraries are synchronous
  and CPU-bound, and a blocking call inside ``async def`` stalls the whole graph
  (050);
* converts driver exceptions into :class:`DocumentParseError` so a corrupt file
  never leaks a ``PdfReadError``/``BadZipFile`` past the infrastructure boundary;
* refuses to guess: an encrypted PDF is rejected instead of being brute-forced;
* records ``truncated`` when a cap is hit rather than dropping text silently (065).

A note on page numbers: PDFs have real pages, so citations map to them. DOCX and
plain text have no pagination primitive, so everything is reported as page 1 —
inventing page numbers would produce citations that cannot be verified.
"""

from __future__ import annotations

import asyncio

from io import BytesIO
from typing import TYPE_CHECKING, Any, ClassVar

from palatium_ai.domain.ports.document_parser import (
    DocumentParseError,
    DocumentParserPort,
    ParsedDocument,
    ParsedPage,
)

if TYPE_CHECKING:
    from collections.abc import Iterable

# Per-page driver failures that mean "this page is unreadable", not "abort the file".
_PAGE_ERRORS: tuple[type[Exception], ...] = (
    ValueError,
    KeyError,
    TypeError,
    AttributeError,
    RecursionError,
)

# Encodings tried in order for text-ish uploads. cp1251 covers Russian CSVs from
# Windows tooling, which is the common real-world case in this deployment.
_TEXT_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "utf-8", "cp1251")


def _truncate_at_word(text: str, max_chars: int) -> str:
    """Prefer a whitespace cut so char caps do not end mid-token."""
    if max_chars <= 0:
        return ""
    if len(text) <= max_chars:
        return text
    piece = text[:max_chars]
    space = piece.rfind(" ")
    if space > max_chars // 2:
        return piece[:space].rstrip()
    return piece.rstrip()


def _decode_text(data: bytes) -> str:
    """Decode bytes with the first encoding that succeeds; never raise on bad bytes."""
    for encoding in _TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


class PlainTextParser:
    """Parser for text/plain, text/markdown and text/csv."""

    def __init__(self, *, media_types: Iterable[str], max_chars: int) -> None:
        self._media_types = frozenset(media_types)
        self._max_chars = max_chars

    def supports(self, mime_type: str) -> bool:
        """Whether this parser handles the media type."""
        return mime_type in self._media_types

    async def parse(self, data: bytes, *, mime_type: str, filename: str) -> ParsedDocument:
        """Decode bytes to text; decoding is CPU-only, so it stays on the loop."""
        text = _decode_text(data)
        truncated = len(text) > self._max_chars
        body = text if not truncated else _truncate_at_word(text, self._max_chars)
        return ParsedDocument(
            pages=(ParsedPage(number=1, text=body),),
            truncated=truncated,
            extraction_source="text_layer",
        )


class PdfParser:
    """Parser for ``application/pdf`` backed by pypdf."""

    _MEDIA_TYPES: ClassVar[frozenset[str]] = frozenset({"application/pdf"})

    def __init__(self, *, max_chars: int, max_pages: int) -> None:
        self._max_chars = max_chars
        self._max_pages = max_pages

    def supports(self, mime_type: str) -> bool:
        """Whether this parser handles the media type."""
        return mime_type in self._MEDIA_TYPES

    async def parse(self, data: bytes, *, mime_type: str, filename: str) -> ParsedDocument:
        """Extract text per page in a worker thread."""
        return await asyncio.to_thread(self._parse_sync, data)

    def _parse_sync(self, data: bytes) -> ParsedDocument:
        try:
            from pypdf import PdfReader
            from pypdf.errors import PdfReadError
        except ImportError as exc:  # pragma: no cover - install-dependent
            msg = "pypdf is required to parse PDFs; install it with `poetry install --with attachments`"
            raise DocumentParseError(msg) from exc

        try:
            reader = PdfReader(BytesIO(data))
        except (PdfReadError, ValueError, OSError) as exc:
            raise DocumentParseError("pdf is corrupt or not a pdf") from exc

        if reader.is_encrypted:
            # Guessing passwords on attacker-supplied input is not a supported flow.
            raise DocumentParseError("pdf is encrypted; decryption is not supported")

        pages: list[ParsedPage] = []
        truncated = False
        total_chars = 0
        for index, page in enumerate(reader.pages, start=1):
            if index > self._max_pages:
                truncated = True
                break
            text = self._extract_page_text(page)
            budget = self._max_chars - total_chars
            if len(text) > budget:
                text = _truncate_at_word(text, max(0, budget))
                truncated = True
            total_chars += len(text)
            pages.append(ParsedPage(number=index, text=text))
            if total_chars >= self._max_chars:
                truncated = True
                break
        return ParsedDocument(
            pages=tuple(pages),
            truncated=truncated,
            extraction_source="text_layer",
        )

    @staticmethod
    def _extract_page_text(page: Any) -> str:
        """One hostile page must not abort the whole document."""
        try:
            return str(page.extract_text() or "")
        except _PAGE_ERRORS:
            return ""


class DocxParser:
    """Parser for ``.docx`` backed by python-docx (paragraphs + table cells)."""

    _MEDIA_TYPES: ClassVar[frozenset[str]] = frozenset(
        {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    )

    #: A ``.docx`` is a zip, and the intake cap bounds the *compressed* size — a
    #: few hundred KB of deflate can inflate to gigabytes of XML. UTF-8 needs at
    #: most 4 bytes per character, so an archive declaring more uncompressed bytes
    #: than the parser could ever turn into characters is a bomb, not a document.
    _MAX_UNCOMPRESSED_BYTES_PER_CHAR: ClassVar[int] = 4
    #: The only entry python-docx turns into paragraphs and table cells.
    _TEXT_PART: ClassVar[str] = "word/document.xml"
    #: Realistic deflate on document XML peaks around 10x; 100x is a bomb.
    _MAX_COMPRESSION_RATIO: ClassVar[int] = 100

    def __init__(self, *, max_chars: int) -> None:
        self._max_chars = max_chars

    def supports(self, mime_type: str) -> bool:
        """Whether this parser handles the media type."""
        return mime_type in self._MEDIA_TYPES

    async def parse(self, data: bytes, *, mime_type: str, filename: str) -> ParsedDocument:
        """Extract text in a worker thread."""
        return await asyncio.to_thread(self._parse_sync, data)

    def _parse_sync(self, data: bytes) -> ParsedDocument:
        from zipfile import BadZipFile

        try:
            from docx import Document
            from docx.opc.exceptions import PackageNotFoundError
        except ImportError as exc:  # pragma: no cover - install-dependent
            msg = "python-docx is required to parse .docx; install it with `poetry install --with attachments`"
            raise DocumentParseError(msg) from exc

        # The zip directory is read first so a decompression bomb is refused
        # before python-docx materialises any XML (020).
        self._assert_archive_within_budget(data)

        try:
            document = Document(BytesIO(data))
        except (PackageNotFoundError, BadZipFile, KeyError, ValueError, OSError) as exc:
            raise DocumentParseError("docx is corrupt or not a docx package") from exc

        blocks = [paragraph.text for paragraph in document.paragraphs]
        blocks.extend("\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        text = "\n".join(block for block in blocks if block.strip())
        truncated = len(text) > self._max_chars
        return ParsedDocument(
            pages=(ParsedPage(number=1, text=text[: self._max_chars]),),
            truncated=truncated,
            extraction_source="text_layer",
        )

    def _assert_archive_within_budget(self, data: bytes) -> None:
        """Refuse a decompression bomb before python-docx materialises any XML.

        Two independent bounds, because the intake cap only sees the *compressed*
        bytes: the text part must fit the character budget, and the archive as a
        whole must not inflate past a realistic deflate ratio. The ratio bound
        leaves room for the boilerplate every docx carries (styles, theme, fonts),
        which a fixed absolute limit would reject for large text budgets.
        """
        from zipfile import BadZipFile, ZipFile

        text_budget = self._max_chars * self._MAX_UNCOMPRESSED_BYTES_PER_CHAR
        total_budget = max(len(data) * self._MAX_COMPRESSION_RATIO, text_budget)
        try:
            with ZipFile(BytesIO(data)) as archive:
                total = 0
                for info in archive.infolist():
                    total += info.file_size
                    if total > total_budget:
                        msg = "docx expands beyond the configured parse budget"
                        raise DocumentParseError(msg)
                    if info.filename == self._TEXT_PART and info.file_size > text_budget:
                        msg = "docx text part exceeds the configured parse budget"
                        raise DocumentParseError(msg)
        except BadZipFile as exc:
            raise DocumentParseError("docx is corrupt or not a docx package") from exc


class XlsxParser:
    """Parser for ``.xlsx`` backed by openpyxl (one :class:`ParsedPage` per sheet)."""

    _MEDIA_TYPES: ClassVar[frozenset[str]] = frozenset(
        {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    )
    _MAX_UNCOMPRESSED_BYTES_PER_CHAR: ClassVar[int] = 4
    _MAX_COMPRESSION_RATIO: ClassVar[int] = 100
    #: Parts openpyxl materialises into cell text during read.
    _TEXT_PART_PREFIXES: ClassVar[tuple[str, ...]] = ("xl/worksheets/", "xl/sharedStrings.xml")

    def __init__(self, *, max_chars: int, max_pages: int) -> None:
        self._max_chars = max_chars
        self._max_pages = max_pages

    def supports(self, mime_type: str) -> bool:
        """Whether this parser handles the media type."""
        return mime_type in self._MEDIA_TYPES

    async def parse(self, data: bytes, *, mime_type: str, filename: str) -> ParsedDocument:
        """Extract sheet text in a worker thread."""
        return await asyncio.to_thread(self._parse_sync, data)

    def _parse_sync(self, data: bytes) -> ParsedDocument:
        from zipfile import BadZipFile

        try:
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException
        except ImportError as exc:  # pragma: no cover - install-dependent
            msg = "openpyxl is required to parse .xlsx; install it with `poetry install --with attachments`"
            raise DocumentParseError(msg) from exc

        self._assert_archive_within_budget(data)

        try:
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, KeyError, ValueError, OSError) as exc:
            raise DocumentParseError("xlsx is corrupt or not an xlsx package") from exc

        pages: list[ParsedPage] = []
        truncated = False
        total_chars = 0
        for sheet_index, sheet_name in enumerate(workbook.sheetnames, start=1):
            if sheet_index > self._max_pages:
                truncated = True
                break
            sheet = workbook[sheet_name]
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) if cell is not None else "" for cell in row]
                line = "\t".join(cells).strip()
                if line:
                    rows.append(line)
            body = "\n".join(rows)
            header = f"Sheet: {sheet_name}\n" if sheet_name else ""
            text = f"{header}{body}".strip()
            budget = self._max_chars - total_chars
            if len(text) > budget:
                text = text[: max(0, budget)]
                truncated = True
            total_chars += len(text)
            pages.append(ParsedPage(number=sheet_index, text=text))
            if total_chars >= self._max_chars:
                truncated = True
                break
        workbook.close()
        if not pages:
            pages.append(ParsedPage(number=1, text=""))
        return ParsedDocument(
            pages=tuple(pages),
            truncated=truncated,
            extraction_source="text_layer",
        )

    def _assert_archive_within_budget(self, data: bytes) -> None:
        from zipfile import BadZipFile, ZipFile

        text_budget = self._max_chars * self._MAX_UNCOMPRESSED_BYTES_PER_CHAR
        total_budget = max(len(data) * self._MAX_COMPRESSION_RATIO, text_budget)
        try:
            with ZipFile(BytesIO(data)) as archive:
                total = 0
                for info in archive.infolist():
                    total += info.file_size
                    if total > total_budget:
                        msg = "xlsx expands beyond the configured parse budget"
                        raise DocumentParseError(msg)
                    name = info.filename.replace("\\", "/")
                    if (
                        any(name.startswith(prefix) for prefix in self._TEXT_PART_PREFIXES)
                        and info.file_size > text_budget
                    ):
                        msg = "xlsx text part exceeds the configured parse budget"
                        raise DocumentParseError(msg)
        except BadZipFile as exc:
            raise DocumentParseError("xlsx is corrupt or not an xlsx package") from exc


class PptxParser:
    """Parser for ``.pptx`` via OOXML zip + lxml (one :class:`ParsedPage` per slide).

    Avoids ``python-pptx`` so the attachments extra stays on already-pinned ``lxml``.
    """

    _MEDIA_TYPES: ClassVar[frozenset[str]] = frozenset(
        {"application/vnd.openxmlformats-officedocument.presentationml.presentation"}
    )
    _MAX_UNCOMPRESSED_BYTES_PER_CHAR: ClassVar[int] = 4
    _MAX_COMPRESSION_RATIO: ClassVar[int] = 100
    _SLIDE_PREFIX: ClassVar[str] = "ppt/slides/slide"

    def __init__(self, *, max_chars: int, max_pages: int) -> None:
        self._max_chars = max_chars
        self._max_pages = max_pages

    def supports(self, mime_type: str) -> bool:
        """Whether this parser handles the media type."""
        return mime_type in self._MEDIA_TYPES

    async def parse(self, data: bytes, *, mime_type: str, filename: str) -> ParsedDocument:
        """Extract slide text in a worker thread."""
        return await asyncio.to_thread(self._parse_sync, data)

    def _parse_sync(self, data: bytes) -> ParsedDocument:
        from zipfile import BadZipFile, ZipFile

        try:
            from lxml import etree
        except ImportError as exc:  # pragma: no cover - install-dependent
            msg = "lxml is required to parse .pptx; install it with `poetry install --with attachments`"
            raise DocumentParseError(msg) from exc

        self._assert_archive_within_budget(data)

        try:
            with ZipFile(BytesIO(data)) as archive:
                slide_names = sorted(
                    name
                    for name in archive.namelist()
                    if name.replace("\\", "/").startswith(self._SLIDE_PREFIX)
                    and name.replace("\\", "/").endswith(".xml")
                )
                pages: list[ParsedPage] = []
                truncated = False
                total_chars = 0
                for slide_index, slide_name in enumerate(slide_names, start=1):
                    if slide_index > self._max_pages:
                        truncated = True
                        break
                    raw = archive.read(slide_name)
                    root = etree.fromstring(raw)
                    texts = [
                        (node.text or "").strip()
                        for node in root.xpath("//*[local-name()='t']")
                        if (node.text or "").strip()
                    ]
                    body = "\n".join(texts)
                    budget = self._max_chars - total_chars
                    if len(body) > budget:
                        body = body[: max(0, budget)]
                        truncated = True
                    total_chars += len(body)
                    pages.append(ParsedPage(number=slide_index, text=body))
                    if total_chars >= self._max_chars:
                        truncated = True
                        break
        except BadZipFile as exc:
            raise DocumentParseError("pptx is corrupt or not a pptx package") from exc
        except etree.XMLSyntaxError as exc:
            raise DocumentParseError("pptx slide XML is corrupt") from exc

        if not pages:
            pages.append(ParsedPage(number=1, text=""))
        return ParsedDocument(
            pages=tuple(pages),
            truncated=truncated,
            extraction_source="text_layer",
        )

    def _assert_archive_within_budget(self, data: bytes) -> None:
        from zipfile import BadZipFile, ZipFile

        text_budget = self._max_chars * self._MAX_UNCOMPRESSED_BYTES_PER_CHAR
        total_budget = max(len(data) * self._MAX_COMPRESSION_RATIO, text_budget)
        try:
            with ZipFile(BytesIO(data)) as archive:
                total = 0
                for info in archive.infolist():
                    total += info.file_size
                    if total > total_budget:
                        msg = "pptx expands beyond the configured parse budget"
                        raise DocumentParseError(msg)
                    name = info.filename.replace("\\", "/")
                    if name.startswith(self._SLIDE_PREFIX) and info.file_size > text_budget:
                        msg = "pptx slide part exceeds the configured parse budget"
                        raise DocumentParseError(msg)
        except BadZipFile as exc:
            raise DocumentParseError("pptx is corrupt or not a pptx package") from exc


class HtmlParser:
    """Parser for ``text/html``: strip scripts/styles, emit visible text (G07)."""

    _MEDIA_TYPES: ClassVar[frozenset[str]] = frozenset({"text/html"})
    _DROP_TAGS: ClassVar[frozenset[str]] = frozenset(
        {"script", "style", "noscript", "template", "iframe", "object", "embed"}
    )

    def __init__(self, *, max_chars: int) -> None:
        self._max_chars = max_chars

    def supports(self, mime_type: str) -> bool:
        """Whether this parser handles the media type."""
        return mime_type in self._MEDIA_TYPES

    async def parse(self, data: bytes, *, mime_type: str, filename: str) -> ParsedDocument:
        """Extract visible text in a worker thread."""
        return await asyncio.to_thread(self._parse_sync, data)

    def _parse_sync(self, data: bytes) -> ParsedDocument:
        try:
            from lxml import (
                etree,
                html as lxml_html,
            )
        except ImportError as exc:  # pragma: no cover - install-dependent
            msg = "lxml is required to parse HTML; install it with `poetry install --with attachments`"
            raise DocumentParseError(msg) from exc

        try:
            document = lxml_html.fromstring(_decode_text(data))
        except (TypeError, ValueError, etree.ParserError) as exc:
            raise DocumentParseError("html is corrupt or not parseable") from exc

        to_drop = []
        for element in document.iter():
            if not isinstance(element.tag, str):
                continue
            local = element.tag.lower().rsplit("}", 1)[-1]
            if local in self._DROP_TAGS:
                to_drop.append(element)
        for element in to_drop:
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)

        text = " ".join(document.text_content().split())
        truncated = len(text) > self._max_chars
        return ParsedDocument(
            pages=(ParsedPage(number=1, text=text[: self._max_chars]),),
            truncated=truncated,
            extraction_source="text_layer",
        )


class CompositeDocumentParser:
    """Dispatches to the first parser that supports the media type."""

    def __init__(self, parsers: Iterable[DocumentParserPort]) -> None:
        self._parsers: tuple[DocumentParserPort, ...] = tuple(parsers)

    def supports(self, mime_type: str) -> bool:
        """Whether any configured parser handles the media type."""
        return any(parser.supports(mime_type) for parser in self._parsers)

    async def parse(self, data: bytes, *, mime_type: str, filename: str) -> ParsedDocument:
        """Parse with the matching parser; unsupported media types fail closed."""
        for parser in self._parsers:
            if parser.supports(mime_type):
                return await parser.parse(data, mime_type=mime_type, filename=filename)
        msg = f"unsupported media type for parsing: {mime_type}"
        raise DocumentParseError(msg)
